from  openpyxl import load_workbook
import os
from config import workbook_name

def is_sheet_present_excel(sheet_name):
    wb=load_workbook(workbook_name)
    if sheet_name in wb.sheetnames:
        return True
    else:
        return False
    
def write_to_cell_excel(worksheet,cell_label,value):
    worksheet[cell_label]=value
    
def get_val_from_excel_file(sheet_name,cell_label):
    wb=load_workbook(workbook_name)
    ws=wb[sheet_name]
    return ws[cell_label].value

def write_in_excel_file(sheet_name,cell_label,value):
    wb=load_workbook(workbook_name)
    ws=wb[sheet_name]
    ws[cell_label]=value
    wb.save(workbook_name)
    
    
    
    