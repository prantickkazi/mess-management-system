from boarder import Boarder 
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from config import (workbook_name,meal_count_sheet_name,marketing_sheet_name,
meal_charge_limit,info_string,success_string,meal_off_string,meal_on_string,
final_balance_info,failure_string,guest_meal_charge,mess_manager_name)

import datetime
import calendar


class Hostel:
    def __init__(self,filename):
        self.filename= filename
        self.boarders= []
        self.manager=mess_manager_name
        
        
    #check status
    def status(self,name):
        wb=load_workbook(workbook_name)
        ws=wb[meal_count_sheet_name]
        now=datetime.datetime.now()
        day=now.day
        present = False
        
        for boarder in self.boarders:
            if boarder.name == name: 
                present = True 
                break
                
        if present:   
            for cell in ws["A"]:
                if cell.value == name:
                    row_number = cell.row
                    break
                
            for cell in ws[1]:
                if cell.value == day:
                    column_no = cell.column
    
            if ws.cell(row=row_number,column=column_no) == meal_on_string or  ws.cell(row=row_number,column=column_no + 2) == meal_on_string:
                ws[f"B{row_number}"] = meal_on_string
            elif ws.cell(row=row_number,column=column_no+1) == meal_on_string or  ws.cell(row=row_number,column=column_no + 3) == meal_on_string:
                ws[f"B{row_number}"] = meal_on_string
            else:
                ws[f"B{row_number}"] = None
        else:
            pass
                
        
    #add boarder and update the excell sheet    
    def add_boarder(self,name):
        boarder=Boarder(name)
        self.boarders.append(boarder)
        wb=load_workbook(workbook_name)
        ws=wb[meal_count_sheet_name]
        last_cell_no=1
        for cl in ws["A"]:
            if cl.value is not None:
                last_cell_no+=1
            else:
                break    
        ws[f"A{last_cell_no}"]=name  
        wb.save(workbook_name)
    
     #add the deposit and updatee the excell sheet   
    def  add_deposit_(self,_name,amount):
        success=False
        target_boarder=None
        for boarder in self.boarders:
            if boarder.name == _name:
                boarder.add_deposit(amount)
                target_boarder=boarder
                success=True
                break
        if success:
            wb=load_workbook(workbook_name)
            ws=wb[meal_count_sheet_name]
            for cell in ws["A"]:
                if cell.value == _name:
                    row_number=cell.row
                    break
            ws[f"C{row_number}"]=target_boarder.deposit
            wb.save(workbook_name)  
            print(f"{success_string}.")
        else:
            print(f"{failure_string}{_name} is not found as a boarder")
     
     #on the meal and update the excell sheet   
    def _meal_on(self,_name):
        success = False
        info = None
        
        #check is user a boarder
        for boarder in self.boarders:
            if boarder.name == _name:
                info=input("""
                "on" for turn on night meal 
                "od" for turn on day meal 
                "yes" for both day and night\n""")
                boarder.on_meal(info)
                success=True
                break
        #update the excell sheet
        if success:
            wb=load_workbook(workbook_name)
            ws=wb[meal_count_sheet_name]
            for cell in ws["A"]:
                if cell.value == _name:
                    row_number=cell.row
                    break
                    
            now=datetime.datetime.now()
            day=now.day 
                  
            #upadate the meal status of particular day
            for cell in ws[1]:
                if cell.value == day:
                    start_col=cell.column
                    break
            if info == "yes":
                if ws.cell(row=row_number,column=start_col).value is None and ws.cell(row=row_number,column=start_col+2).value is None:
                    ws.cell(row=row_number,column=start_col).value = meal_on_string
                    ws.cell(row=row_number,column=start_col+2).value= meal_on_string
                    ws[f"B{row_number}"]=meal_on_string
                    if ws[f"D{row_number}"].value is None:
                        ws[f"D{row_number}"].value = 0
                        ws[f"D{row_number}"].value +=2
                    else:
                        ws[f"D{row_number}"].value +=2
                    print(f"{success_string}")
        
            elif info =="od":
                if ws.cell(row=row_number,column=start_col).value is None :
                    ws.cell(row=row_number,column=start_col).value = meal_on_string
                    ws[f"B{row_number}"]=meal_on_string
                    if ws[f"D{row_number}"].value is None:
                        ws[f"D{row_number}"].value = 0
                        ws[f"D{row_number}"].value +=1
                    else:
                        ws[f"D{row_number}"].value +=1
                print({success_string})
                
            elif info =="on":
                if ws.cell(row=row_number,column = start_col+2).value is None:
                    ws.cell(row=row_number,column = start_col+2).value= meal_on_string
                    ws[f"B{row_number}"] = meal_on_string
                    if ws[f"D{row_number}"].value is None:
                        ws[f"D{row_number}"].value = 0
                        ws[f"D{row_number}"].value +=1
                    else:
                        ws[f"D{row_number}"].value +=1
                    print(success_string)
            else:
                print(f"{failure_string} invalid responce")
                    
            self.status(_name)
            wb.save(workbook_name)
            
        else:
            print(f"Name {_name} does not exist {failure_string}")
            
    #off the meal and update the excell sheet
    def _meal_off(self,_name):
        info=None
        success=False
        #check in memory if user exist or not
        for boarder in self.boarders:
            if boarder.name == _name:
                info=input("""
                "on" for turn off night meal 
                "od" for turn off day meal 
                "off" for both day and night\n""")
                boarder.off_meal(info)
                success=True
        
        if success:       
            #update the excel sheet
            wb=load_workbook(workbook_name)
            ws=wb[meal_count_sheet_name]
            now=datetime.datetime.now()
            date=now.day
            
            for cell in ws["A"]:
                if cell.value == _name:
                    row_number = cell.row
                    break
            for cell in ws[1]:
                if cell.value == date:
                    start_col=cell.column
                    break
            if info == "off":
                if  ws.cell(row=row_number,column=start_col).value == meal_on_string and  ws.cell(row=row_number,column=start_col+2).value == meal_on_string:
                    ws.cell(row=row_number,column=start_col).value = None
                    ws.cell(row=row_number,column=start_col+2).value = None 
                    ws[f"B{row_number}"].value = None
                    ws[f"D{row_number}"].value -= 2
                    print(success_string)
            elif info == "od":
                if ws.cell(row=row_number,column=start_col).value == meal_on_string:
                    ws.cell(row=row_number,column=start_col).value = None
                    ws[f"D{row_number}"].value -= 1
                    print(success_string)
            elif info == "on":
                if  ws.cell(row=row_number,column=start_col+2).value == meal_on_string:
                    ws.cell(row=row_number,column=start_col+2).value = None
                    ws[f"D{row_number}"].value -= 1
                    print(success_string)
            else:
                print(f"{failure_string} invalid responce")
            
            self.status(_name)
            wb.save(workbook_name)
            
        else:
            print(f"{_name} you are not a boarder {failure_string}")
    
    #guest meal on 
    def guest_meal_on(self,_name):
        success=False
        string_msg= True
        for object in self.boarders:
            if object.name == _name:
                string_msg=False
                info=input("""
                on guest meal in both day and night -yes
                on guest meal in day-od
                on guest meal at night-on\n
                           """)
                if object.day_active or object.night_active :
                    object.on_guest_meal(info)
                    success=True
                    break
                else:
                    print(f"{failure_string} you dont have any active meal status")
                    break
        if success:
            wb=load_workbook(workbook_name)
            ws=wb[meal_count_sheet_name]
            now=datetime.datetime.now()
            day=now.day
            for cell in ws["A"]:
                if cell.value == _name:
                    row_number=cell.row
                    break
            for cell in ws[1]:
                if cell.value == day:
                    get_col=cell.column
                    break
            if info == "od":
                if ws.cell(row=row_number,column=get_col+1).value is None:
                    ws.cell(row=row_number,column=get_col+1).value=meal_on_string
                    if ws[f"E{row_number}"].value is None:
                        ws[f"E{row_number}"].value = 0
                        ws[f"E{row_number}"].value +=1
                        print(success_string)
                    else:
                        ws[f"E{row_number}"].value +=1
                        print(success_string)
            elif info == "on":
                if ws.cell(row=row_number,column=get_col+3).value is None:
                   ws.cell(row=row_number,column=get_col+3).value = meal_on_string
                   if ws[f"E{row_number}"].value is  None:
                        ws[f"E{row_number}"].value = 0
                        ws[f"E{row_number}"].value +=1
                        print(success_string)
                   else:
                        ws[f"E{row_number}"].value +=1
                        print(success_string)
            elif info == "yes":
                if ws.cell(row=row_number,column=get_col+1).value is None and  ws.cell(row=row_number,column=get_col+3).value is None:
                    ws.cell(row=row_number,column=get_col+1).value = meal_on_string
                    ws.cell(row=row_number,column=get_col+3).value = meal_on_string
                    if ws[f"E{row_number}"].value is None:
                        ws[f"E{row_number}"].value = 0
                        ws[f"E{row_number}"].value +=2
                        print(success_string)
                    else:
                        ws[f"E{row_number}"].value +=2
                        print(success_string)
            else:
                print(f"{failure_string} invalid responce")
            self.status(_name)   
            wb.save(workbook_name)
            
        elif string_msg:
             print(f"no user exist wuth this name {failure_string}") 
            
    def guest_meal_off(self,_name):
        success=False
        for object in self.boarders:
            if object.name == _name:
                info=input("""
                off guest meal in both day and night -off
                off guest meal in day-od
                off guest meal at night-on\n
                           """)
                object.off_guest_meal(info)
                success=True
                break
        if success:
            wb=load_workbook(workbook_name)
            ws=wb[meal_count_sheet_name]
            now=datetime.datetime.now()
            day=now.day
            for cell in ws["A"]:
                if cell.value == _name:
                    row_number=cell.row
                    break
            for cell in ws[1]:
                if cell.value == day:
                    get_col=cell.column
                    break
            if info == "od":
                if ws.cell(row=row_number,column=get_col+1).value == meal_on_string:
                    ws.cell(row=row_number,column=get_col+1).value=None
                    ws[f"E{row_number}"].value -= 1
                    print(success_string)
            elif info == "on":
                if ws.cell(row=row_number,column=get_col+3).value == meal_on_string:
                    ws.cell(row=row_number,column=get_col+3).value=None
                    ws[f"E{row_number}"].value -= 1
                    print(success_string)
            elif info == "off":
                if  ws.cell(row=row_number,column=get_col+1).value == meal_on_string and ws.cell(row=row_number,column=get_col+3).value == meal_on_string :
                    ws.cell(row=row_number,column=get_col+1).value=None
                    ws.cell(row=row_number,column=get_col+3).value=None
                    ws[f"E{row_number}"].value -= 2
                    print(success_string)
            else:
                print(f"{failure_string} invalid responce")
            
            self.status(_name) 
            wb.save(workbook_name)
            
        else:
            print(f"no user exist wuth this name {failure_string}") 
            
    
    def total_expences(self):
        wb=load_workbook(workbook_name)
        ws=wb[marketing_sheet_name]
        now=datetime.datetime.now()
        year=now.year
        month=now.month
        day=calendar.monthrange(year=year,month=month)[1]
        daily_marketing_expences=0
        #daily marketing expences
        for row in range (3,day+3):
            if ws.cell(row=row,column=3).value is None:
                ws.cell(row=row,column=3).value = 0
                daily_marketing_expences += ws.cell(row=row,column=3).value
            else:
                daily_marketing_expences += int(ws.cell(row=row,column=3).value)
        #total expences
        total=daily_marketing_expences 
        return total
                
    def guest_meal_counts(self):
        total_boarder=len(self.boarders)
        count=[]
        loop_count=0
        wb=load_workbook(workbook_name)
        ws=wb[meal_count_sheet_name]
        for cell in ws["E"]:
            if loop_count <= total_boarder + 2:
                if cell.value is not None:
                    count.append(cell.value)
                else:
                    count.append(0)
                loop_count += 1
            else:
                break
        modified_count=count[2:]
        return sum(modified_count)
        
    def own_meal_counts(self):
        total_boarders=len(self.boarders)
        loop_count=0
        count=[]
        wb=load_workbook(workbook_name)
        ws=wb[meal_count_sheet_name]
        for cell in ws["D"]:
            if loop_count <= total_boarders +2:
                if cell.value is not None:
                    count.append(cell.value)
                else:
                    count.append(0)
                loop_count+=1
            else:
                break
        updated_count=count[2:]
        return sum(updated_count)
    
    def hostel_mess_balance(self):
        deposit_balance=0
        expences=self.total_expences()
        for object in self.boarders:
            deposit_balance+= object.deposit
        current_balance=deposit_balance - expences
        if current_balance < 0:
            print(f"{-1*current_balance} due in the market from the mess" )
        elif current_balance >= 0:
            print(f"Balance {current_balance}")
        print(f"Total deposit :{deposit_balance},   Total expences: {expences}")
        return current_balance
               
    def calculate_meal_charge(self):
        expences=self.total_expences()
        all_meals=self.guest_meal_counts() + self.own_meal_counts()
        if expences == 0 or all_meals == 0:
            print("Tere is no expences or no meal is active in past")
        else: 
            if all_meals == 0:
                print("Mess just starts today")
                meal_charge=0
            else:
                meal_charge = expences / all_meals
        info_list=[meal_charge,expences,all_meals]
        return  info_list
        
    def update_daily_marketing(self):
        wb=load_workbook(workbook_name)
        ws=wb[marketing_sheet_name]
        now=datetime.datetime.now()
        year=now.year
        month=now.month
        day=now.day
        for cell in ws["A"]:
            if cell.value == f"{day}/{month}/{year}":
                row_num=cell.row
                break
       
        ws.cell(row=row_num,column=2).value=grocery_shoper=input("Enter the grocery provider's name --- ")
        ws.cell(row=row_num,column=3).value=money_spend=input("Enter the money spent in daily marketing ---")
        ws.cell(row=row_num,column=4).value=remark=input("Enter the remarks ---")
        print(success_string)
        wb.save(workbook_name)
    
    def txtfile_final_balance(self):
        balance_dict={}
        if len(self.boarders) > 0:
            for boarder_obb in self.boarders:
                boarder_final_balance_info=[]
                own_meal_expences = (boarder_obb.total_own_meal * self.calculate_meal_charge()[0])
                guest_meal_expences=(boarder_obb.total_guest_meal * guest_meal_charge)
                boarder_final_balance_info.append(own_meal_expences + guest_meal_expences)
                balance=boarder_obb.deposit - (own_meal_expences + guest_meal_expences)
                boarder_final_balance_info.append(balance)
                balance_dict[boarder_obb.name]=boarder_final_balance_info
            with open(final_balance_info,"w") as file:
                sl_no=1
                for key,value in balance_dict.items():
                    if value[1] == 0:
                        file.write(f"{sl_no}. {key}----Your money in current mess is clear\n")
                    elif value[1] < 0:
                        file.write(f"{sl_no}. {key}----Money due {round(-1*value[1],2)}\n")
                    elif value[1] > 0:
                        file.write(f"{sl_no}. {key}----return from mess {round(value[1],2)}\n")
                    sl_no += 1
                    
        else:
            print("No boarder exist")
                    
    def daily_number_of_meals(self):
        day = 0
        guest_day = 0
        night = 0
        guest_night =0
        now=datetime.datetime.now()
        date=now.day
        loop=len(self.boarders)        
        wb=load_workbook(workbook_name)
        ws=wb[meal_count_sheet_name]
        for cell in ws[1]:
            if cell.value == date:
                colum_no=cell.column
        for row in range(3,loop + 3):
            cell= ws.cell(row=row,column=colum_no)
            if cell.value == meal_on_string:
                day += 1
        
        for row in range(3,loop + 3):
            cell= ws.cell(row=row,column=colum_no + 1)
            if cell.value == meal_on_string:
                guest_day += 1
                
        for row in range(3,loop + 3):
            cell= ws.cell(row=row,column=colum_no + 2)
            if cell.value == meal_on_string:
                night += 1
        
        for row in range(3,loop + 3):
            cell= ws.cell(row=row,column=colum_no + 3)
            if cell.value == meal_on_string:
                guest_night += 1
                
        info_tup=(day,guest_day,night,guest_night)
        info_list=list(info_tup)
        
        return info_list
        
    def create_boarders_txt(self):
        with open("boarder.txt","w") as file:
            sl_no=1
            if self.boarders:
                for boarer_obb in self.boarders:
                    file.write(f"{sl_no}.{boarer_obb.name}\n")
                    sl_no += 1
            else:
                print("No boarders exist in the hostel ")
                
    def add_boarder_from_txt(self):
        name_list=[]
        for obb in self.boarders:
            name_list.append(obb.name)
        with open("boarder.txt", "r") as file:
            for name in file:
                clean_line=name.rstrip("\n")
                if clean_line[2:] not in name_list:
                    self.add_boarder(clean_line[2:])
        
       
        
        
        
    
            
        
        
        
            
            
                    
                    
                        
            
            

   
            
            
      
        