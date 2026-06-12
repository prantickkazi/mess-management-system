from openpyxl import load_workbook
from meal_count_template import initialize_meal_count_sheet,create_meal_count_sheet
from meal_routin_template import init_meal_routin,create_meal_routin_sheet
from marketing_template import init_marketing_sheet,create_marketing_sheet
from boarder import Boarder
from hostel import Hostel
from config import (workbook_name,meal_count_sheet_name,
info_string,success_string,meal_off_string,meal_on_string,meal_charge_limit,failure_string)

import datetime
import calendar
import sys
import os

def create_all_sheets():
     initialize_meal_count_sheet()
     create_meal_count_sheet()
     init_meal_routin()
     create_meal_routin_sheet()
     init_marketing_sheet()
     create_marketing_sheet()
     


def day_month_year():
     now=datetime.datetime.now()
     month=now.month
     year=now.year
     day=now.day
     date_list=[day,month,year]
     return date_list       

def load_from_database(hostel_object):
    loaded=False
    wb=load_workbook(workbook_name)
    ws=wb[meal_count_sheet_name]
    now=datetime.datetime.now()
    date=now.day
    for cell in ws["A"]:
        if cell.value is None:
            break
        else:
            boarder_obb=Boarder(cell.value)
            row_no=cell.row 
            for cell in ws[1]:
                if cell.value == date:
                    start_col=cell.column
                    break
            #load boarder attributes 
            boarder_obb.row_number=row_no

            boarder_obb.deposit= 0 if ws[f"C{row_no}"].value is None else ws[f"C{row_no}"].value
            boarder_obb.guest_active= True if ws.cell(row=row_no, column=start_col+1).value == meal_on_string or ws.cell(row=row_no,column=start_col+3).value == meal_on_string else None
            boarder_obb.active=True if  ws.cell(row=row_no, column=start_col).value == meal_on_string or ws.cell(row=row_no,column=start_col+2).value == meal_on_string  else None 
            boarder_obb.day_active= True if ws.cell(row=row_no, column=start_col ).value == meal_on_string else None
            boarder_obb.night_active= True if ws.cell(row=row_no, column=start_col + 2).value == meal_on_string else None
            boarder_obb.guest_day_active=True if ws.cell(row=row_no, column=start_col+1).value == meal_on_string  else None 
            boarder_obb.guest_night_active=True if ws.cell(row=row_no, column= start_col + 3) == meal_on_string  else None
            boarder_obb.total_guest_meal= 0 if  ws[f"E{row_no}"].value == None else ws[f"E{row_no}"].value
            boarder_obb.total_own_meal= 0 if ws[f"D{row_no}"].value == None else ws[f"D{row_no}"].value
            hostel_object.boarders.append(boarder_obb)
            loaded=True
    if loaded:
            del hostel_object.boarders[0:2]
            print(f"no of all boarders --{len(hostel_object.boarders)}")
     
    return hostel_object.boarders

def own_work(hostel_object):
     with open("boarder.txt") as file:
          for name in file:
               clean_name=name.rstrip("\n")
               print(f"{clean_name}")
     user_input=input("Enter the name of user --")
     user_name=str(user_input)
     wb=load_workbook(workbook_name)
     ws=wb[meal_count_sheet_name]
     user_present=False
     for cell in ws["A"]:
          if cell.value == user_name:
               user_present=True
               break
     if user_present:
          print(success_string)
          while user_present:
               print(f"{info_string}-Enter 00 to exit  and {info_string} r to refresh") 
               print("--------------welcome to your profile---------------")
               user_input=input("""
               0.Meal Statues               1.Meal On
               2.Meal off                   3.Guest Meal on
               4.Guest Meal off             5.Deposit amount
               6.Current Balance            7.Total Meals\n
               """)
               if user_input == "0":
                    meal_type=input("""
                                   1.Own meal
                                   2.Guest meal
                    """)
                    if meal_type =="1":
                         for boarder in hostel_object.boarders:
                              if boarder.name == user_name:
                                   status=boarder.active
                                   if status:
                                        print(f"{user_name } your meal is active {meal_on_string}")
                                   else:
                                        print(f"{user_name} your meal is not active {meal_off_string}")
                                   pass
                    elif meal_type == "2":
                         for boarder in hostel_object.boarders:
                              if boarder.name == user_name:
                                   status= boarder.guest_active
                                   if status:
                                        print(f"{user_name} your guest meal is active {meal_on_string}")
                                   else:
                                        print(f"{user_name} your guest meal is not active {meal_off_string}")
                         pass
               elif user_input == "1":
                    hostel_object._meal_on(user_name)
               elif user_input == "2":
                    hostel_object._meal_off(user_name)
               elif user_input == "3":
                    hostel_object.guest_meal_on(user_name)
               elif user_input == "4":
                    hostel_object.guest_meal_off(user_name)
               elif user_input == "5":
                    deposit_amount=int(input("enter the deposit amount only intiger value---"))
                    hostel_object.add_deposit_(user_name,deposit_amount)
               elif user_input == "6":
                    boarers_list=hostel_object.boarders
                    for object in boarers_list:
                         if object.name == user_name:
                              balance=object.deposit
                              print(f"{info_string}{user_name} your balance in mess is -- {balance}")
               elif user_input =="7":
                    for object in hostel_object.boarders:
                         if object.name == user_name:
                              print(f"your total meal till now : {object.total_own_meal + object.total_guest_meal}")
               elif user_input == "00":
                    break
               else:
                    print(f"--{info_string}")
     else:
          print("User not exist",failure_string)
               
def main_menu():
     my_hostel=Hostel(workbook_name)
     load_from_database(my_hostel)
     my_hostel.create_boarders_txt()
     my_hostel.add_boarder_from_txt()
     print(f"{info_string}-Enter 00 to exit  and {info_string} r to refresh") 
     user_responce=input("""
                  -----------Menu----------\n
     ------enter the corrosponding serial number -------\n
     1.User                                 2.Add new Boarder
     3.Calculate Meal charge                4.Update daily marketing               
     5.Current mess balance                 6.Final Balance Sheet                  
     7.Number of meals                      8.Food menu\n""")
     my_hostel.add_boarder_from_txt()
     if user_responce == "1":
          own_work(my_hostel)
     if user_responce == "2":
          name=input("enter the boarder name: ")
          my_hostel.add_boarder(name=name)
          print(f"{name}: boarder added to to hostel mess {success_string}")
     elif user_responce == "3":
          meal_charge = my_hostel.calculate_meal_charge()[0]
          print(f"Meal charge: {round(meal_charge,2)}")
          print(f"Total expence : {round(my_hostel.calculate_meal_charge()[1],2)}")
          print(f"Total meals : {round(my_hostel.calculate_meal_charge()[2],2)}")
          if meal_charge <= meal_charge_limit :
               print(f"Meal_charge is under current meal charge limit")
          else:
               print(f"Meal charge is above current meal charge limit")
     elif user_responce == "4":
          my_hostel.update_daily_marketing()
     elif user_responce == "00":
          sys.exit()
     elif user_responce == "5":
          my_hostel.hostel_mess_balance()
     elif user_responce == "6":
          my_hostel.txtfile_final_balance()
     elif user_responce == "7":
          information = my_hostel.daily_number_of_meals()
          print(f" Day --{information[0]}\n",f"Guest Day --{information[1]}\n",f"Night --{information[2]}\n",f"Guest Night --{information[3]}\n")
     elif user_responce == "8":
          print("not developed yet")
     else:
          print(f"---{info_string}")
