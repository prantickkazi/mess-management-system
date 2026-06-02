from openpyxl import load_workbook
from openpyxl.styles import Font,Alignment,PatternFill
from openpyxl.utils import get_column_letter
from datetime import timedelta

import datetime
import calendar

from config import workbook_name,marketing_sheet_name

#init marketing sheet
def init_marketing_sheet():
    wb=load_workbook(workbook_name)
    ws=wb.create_sheet(marketing_sheet_name)
    wb.save(workbook_name)
#create marketing sheet
def create_marketing_sheet():
    wb=load_workbook(workbook_name)
    ws=wb[marketing_sheet_name]
    
    #get date and year
    now=datetime.datetime.now()
    month=now.month
    year=now.year
    days_in_month=calendar.monthrange(year,month)[1]
    month_name=now.strftime("%B, %Y")
    
    #creating the rows
    row1=[f"Daily Marketing {month_name}", "", "", "",]
    ws.append(row1)
    row2=["Date","Name","Money Spent","Remark"]
    ws.append(row2)
    d=1
    for row in range(3,days_in_month+3):
        ws.cell(row=row,column=1).value=f"{d}/{month}/{year}"
        d+=1
    
    
    #styling
    #colour and font
    black_bolt=Font(name="Calibri",size=11,bold=True,color="000000")
    center_align=Alignment(vertical="center",horizontal="center")
    heaader_fill=PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
    
    #applying it to cells
    for row in range(1,3):
        for col in range(1,5):
            cell=ws.cell(row=row,column=col)
            cell.font=black_bolt
            cell.fill=heaader_fill
            cell.alignment=center_align
    
    #formating and merging
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=4)
    
    for col in range(1,5):
        col_lebel=get_column_letter(col)
        ws.column_dimensions[col_lebel].width=15
    wb.save(workbook_name)
    print(f"Sheet {marketing_sheet_name} is successfully created in {workbook_name}")
    
# init_marketing_sheet()
# create_marketing_sheet()
       
            
    
    
    
    
    