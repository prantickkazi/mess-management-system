from openpyxl import workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from config import meal_routin_sheet_name,workbook_name

import datetime
import calendar

#initilize the sheet
def init_meal_routin():
    wb=load_workbook(workbook_name)
    ws=wb.create_sheet(meal_routin_sheet_name)
    wb.save(workbook_name)
    
#cereate the sheet
def create_meal_routin_sheet():
    wb=load_workbook(workbook_name)
    ws=wb[meal_routin_sheet_name]
    row1=["Date","Veg","Egg","Fish","Chicken Curry","Other"]
    ws.append(row1)
    
    #month and year
    now=datetime.datetime.now()
    month=now.month
    year=now.year
    days_in_month=calendar.monthrange(year=year,month=month)[1]
    d=1
    for date in range(2,days_in_month+2):
        cell=ws.cell(row=date,column=1).value=f"{d}/{month}/{year}"
        d+=1
        
    #styling
    black_bolt=Font(name="Calibri",size=11,bold=True,color="000000")
    center_align=Alignment(horizontal="center",vertical="center")
    veg_header=PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    egg_header=PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fish_header=PatternFill(start_color="0000FF", end_color="0000FF", fill_type="solid")
    meat_header=PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid')
    other_header=PatternFill(start_color="FFD1DC", end_color="FFD1DC", fill_type="solid")
    
    for c in ws[1]:
        c.font=black_bolt
        c.alignment=center_align
    
    ws.cell(row=1,column=2).fill=veg_header
    ws.cell(row=1,column=3).fill=egg_header
    ws.cell(row=1,column=4).fill=fish_header
    ws.cell(row=1,column=5).fill=meat_header
    ws.cell(row=1,column=6).fill=other_header
            
    #formating
    for c in range(1,7):
        col_letter=get_column_letter(c)
        ws.column_dimensions[col_letter].width=16
        
    wb.save(workbook_name)
    print(f"Sheet {meal_routin_sheet_name} is successfully created in {workbook_name}")
    
# init_meal_routin()
# create_meal_routin_sheet()
    
        